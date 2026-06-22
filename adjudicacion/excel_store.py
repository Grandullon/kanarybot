"""
Gestión de los Excel para el acto de adjudicación de contratos.

Trabaja con DOS archivos (o dos hojas), y CADA aceptación se refleja en LOS DOS:

  1. CANDIDATOS  (listado definitivo de inscritos)
     - El Nº de candidato = su posición en la lista (1, 2, 3, ...): es lo que se canta
       en el acto ("el 177, fulanito, acepta este contrato").
     - Al asignar/cambiar estado se escribe en sus columnas de oferta
       (Oferta Estado / Oferta Aceptada / Oferta Datos): así queda constancia de que esa
       persona YA tiene contrato aceptado.

  2. PUESTOS / SELECCIÓN  (relación de contratos a ofertar)
     - Columnas tipo: centro, fecha inicio, fecha fin, duracion, turno, seleccinado,
       estado, anotacion.
     - Al asignar se escribe el nombre del candidato en «seleccinado» y el estado. En cuanto
       el estado es «Aceptado», el puesto deja de contar como "por ofertar".

Reglas de visualización:
  - Un puesto está ADJUDICADO (deja de ofertarse) cuando su estado es «Aceptado».
  - Cualquier otro estado (Pendiente, Contactado, No contesta, Renuncia) lo mantiene en la
    lista de "por ofertar" (sigue disponible / en gestión).

Seguridad de escritura: copia de seguridad previa + guardado atómico, todo bajo un único
bloqueo (un solo operador/proceso). Si un Excel está abierto, se avisa y no se pierde nada.
"""

import os
import sys
import json
import shutil
import tempfile
import threading
import unicodedata
from datetime import datetime

from openpyxl import load_workbook

ESTADOS = ["Aceptado", "Contactado", "No contesta", "Renuncia", "Pendiente"]
ESTADO_INICIAL = "Pendiente"
ESTADO_ADJUDICADO = "Aceptado"  # único estado que retira el puesto de "por ofertar"

# Columnas auxiliares que se añaden al archivo de puestos si no existen ya.
COL_NUM = "Nº candidato"
COL_DNI_AUX = "DNI candidato"
COL_HORA = "Hora"

# Filas vacías seguidas que marcan el final de los datos (Excel reporta 1.048.576 filas).
_RACHA_VACIAS = 80


# --------------------------------------------------------------------- #
#  Utilidades
# --------------------------------------------------------------------- #

def _normaliza(texto):
    """Minúsculas sin acentos ni espacios sobrantes, para comparar y buscar."""
    if texto is None:
        return ""
    txt = unicodedata.normalize("NFKD", str(texto))
    txt = "".join(c for c in txt if not unicodedata.combining(c))
    return " ".join(txt.split()).lower()


def _adivina(cabeceras, palabras_clave):
    """Primera cabecera cuyo texto normalizado contiene alguna palabra clave."""
    for cab in cabeceras:
        norm = _normaliza(cab)
        for clave in palabras_clave:
            if clave in norm:
                return cab
    return None


def _valor(fila, idx):
    """Valor limpio de una celda; '' si no existe. Enteros sin '.0'; fechas dd/mm/aaaa."""
    if idx is None or fila is None or idx >= len(fila):
        return ""
    v = fila[idx]
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return v.strip() if isinstance(v, str) else str(v)


def _cabeceras(ws):
    for fila in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        return [c if c is not None else "" for c in fila]
    return []


def _idx(cabeceras, nombre):
    """Índice 0-based de una columna por su cabecera (comparación normalizada)."""
    if not nombre:
        return None
    objetivo = _normaliza(nombre)
    for i, cab in enumerate(cabeceras):
        if _normaliza(cab) == objetivo:
            return i
    return None


def _escribe(ws, fila, cab, nombre_col, valor):
    """Escribe `valor` en la celda (fila, columna-por-cabecera) si la columna existe."""
    i = _idx(cab, nombre_col)
    if i is not None:
        ws.cell(row=fila, column=i + 1).value = valor


def _filas_datos(ws):
    """Genera (numero_fila, tupla_valores) saltando filas vacías; corta tras una racha."""
    vacias = 0
    for nfila, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if fila is None or all(v is None or str(v).strip() == "" for v in fila):
            vacias += 1
            if vacias >= _RACHA_VACIAS:
                break
            continue
        vacias = 0
        yield nfila, fila


# --------------------------------------------------------------------- #
#  Almacén
# --------------------------------------------------------------------- #

class ExcelStore:
    """Acceso a los Excel de candidatos y de puestos. Toda escritura va bajo bloqueo."""

    def __init__(self, config):
        self.config = config
        self.lock = threading.Lock()
        self.ultimo_error = None
        self._cache_candidatos = None  # el listado no cambia durante el acto

    # ----------------------- Auto-detección (/config) ----------------- #

    @staticmethod
    def info_excel(ruta):
        wb = load_workbook(ruta, read_only=True, data_only=True)
        info = {"hojas": wb.sheetnames, "cabeceras": {}}
        for h in wb.sheetnames:
            info["cabeceras"][h] = _cabeceras(wb[h])
        wb.close()
        return info

    @staticmethod
    def autodetectar_candidatos(ruta, hoja=None):
        wb = load_workbook(ruta, read_only=True, data_only=True)
        hoja = hoja or wb.sheetnames[0]
        cab = _cabeceras(wb[hoja])
        wb.close()
        return {
            "ruta": ruta, "hoja": hoja,
            "col_nombre": _adivina(cab, ["nombre"]) or "",
            "col_apellidos": _adivina(cab, ["apellido"]) or "",
            "col_dni": _adivina(cab, ["dni", "nif", "documento"]) or "",
            "col_letra": _adivina(cab, ["letra"]) or "",
            "col_telefono": _adivina(cab, ["tlf", "telefono", "movil"]) or "",
            "col_puntos": _adivina(cab, ["autobaremo", "total puntos", "puntos"]) or "",
            "col_oferta_estado": _adivina(cab, ["oferta estado", "estado oferta"]) or "",
            "col_oferta_aceptada": _adivina(cab, ["oferta aceptada", "aceptada"]) or "",
            "col_oferta_datos": _adivina(cab, ["oferta datos", "datos oferta"]) or "",
            "cabeceras": cab,
        }

    @staticmethod
    def autodetectar_puestos(ruta, hoja=None):
        wb = load_workbook(ruta, read_only=True, data_only=True)
        hoja = hoja or wb.sheetnames[0]
        cab = _cabeceras(wb[hoja])
        wb.close()
        return {
            "ruta": ruta, "hoja": hoja,
            "col_centro": _adivina(cab, ["centro", "hospital", "distrito", "destino"]) or "",
            "col_fecha_inicio": _adivina(cab, ["fecha inicio", "inicio", "alta"]) or "",
            "col_fecha_fin": _adivina(cab, ["fecha fin", "fin", "hasta"]) or "",
            "col_duracion": _adivina(cab, ["duracion", "dias", "meses", "tiempo", "periodo"]) or "",
            "col_turno": _adivina(cab, ["turno", "jornada"]) or "",
            "col_nombre_acepta": _adivina(cab, ["seleccin", "seleccion", "selecc", "profesional acepta", "nombre acepta", "asignado", "acepta"]) or "",
            "col_dni_acepta": _adivina(cab, ["dni acepta", "dni profesional", "dni asignado", "dni"]) or "",
            "col_estado": _adivina(cab, ["estado"]) or "",
            "col_anotacion": _adivina(cab, ["anotac", "observ", "nota"]) or "",
            "cabeceras": cab,
        }

    # ----------------------- Candidatos (lectura) --------------------- #

    def candidatos(self):
        if self._cache_candidatos is not None:
            return self._cache_candidatos

        c = self.config["candidatos"]
        wb = load_workbook(c["ruta"], read_only=True, data_only=True)
        try:
            ws = wb[c["hoja"]]
            cab = _cabeceras(ws)
            i_nom = _idx(cab, c.get("col_nombre"))
            i_ape = _idx(cab, c.get("col_apellidos"))
            i_dni = _idx(cab, c.get("col_dni"))
            i_let = _idx(cab, c.get("col_letra"))
            i_tlf = _idx(cab, c.get("col_telefono"))
            i_pts = _idx(cab, c.get("col_puntos"))

            lista = []
            numero = 0
            for nfila, fila in _filas_datos(ws):
                numero += 1
                nombre = _valor(fila, i_nom)
                apellidos = _valor(fila, i_ape)
                lista.append({
                    "numero": numero,       # posición en el listado definitivo
                    "fila": nfila,          # fila real en el Excel (para escribir de vuelta)
                    "nombre": (nombre + " " + apellidos).strip(),
                    "nombre_pila": nombre,  # solo el nombre (sin apellidos), para la pantalla
                    "dni": self._dni_completo(_valor(fila, i_dni), _valor(fila, i_let)),
                    "telefono": _valor(fila, i_tlf),
                    "puntos": _valor(fila, i_pts),
                })
        finally:
            wb.close()

        self._cache_candidatos = lista
        return lista

    @staticmethod
    def _dni_completo(numero, letra):
        numero = (numero or "").replace(" ", "")
        letra = (letra or "").strip()
        return (numero + letra).strip()

    def candidato_por_numero(self, numero):
        try:
            numero = int(numero)
        except (TypeError, ValueError):
            return None
        for cand in self.candidatos():
            if cand["numero"] == numero:
                return cand
        return None

    def buscar_candidatos(self, consulta, limite=15):
        """Busca por número (exacto/prefijo), nombre o DNI, ordenando por relevancia.

        Prioridad: nº exacto > nº que empieza por > nombre > DNI. Así, al cantar
        "el 177", el candidato nº 177 aparece el primero aunque haya DNIs con un 177.
        """
        q = _normaliza(consulta)
        if not q:
            return []
        puntuados = []
        for cand in self.candidatos():
            num = str(cand["numero"])
            if q == num:
                score = 0
            elif num.startswith(q):
                score = 1
            elif q in _normaliza(cand["nombre"]):
                score = 2
            elif q in _normaliza(cand["dni"]):
                score = 3
            else:
                continue
            puntuados.append((score, cand["numero"], cand))
        puntuados.sort(key=lambda x: (x[0], x[1]))
        return [c for _, _, c in puntuados[:limite]]

    # ----------------------- Puestos (lectura) ------------------------ #

    def puestos(self):
        p = self.config["puestos"]
        wb = load_workbook(p["ruta"], data_only=True)
        try:
            ws = wb[p["hoja"]]
            cab = _cabeceras(ws)
            i = {
                "centro": _idx(cab, p.get("col_centro")),
                "fini": _idx(cab, p.get("col_fecha_inicio")),
                "ffin": _idx(cab, p.get("col_fecha_fin")),
                "dur": _idx(cab, p.get("col_duracion")),
                "turno": _idx(cab, p.get("col_turno")),
                "nom": _idx(cab, p.get("col_nombre_acepta")),
                "dni": _idx(cab, p.get("col_dni_acepta")),
                "est": _idx(cab, p.get("col_estado")),
                "anot": _idx(cab, p.get("col_anotacion")),
                "numc": _idx(cab, COL_NUM),
                "dniaux": _idx(cab, COL_DNI_AUX),
                "hora": _idx(cab, COL_HORA),
            }
            lista = []
            for nfila, fila in _filas_datos(ws):
                estado = _valor(fila, i["est"]) or ESTADO_INICIAL
                asignado = _valor(fila, i["nom"])
                lista.append({
                    "id": nfila,
                    "centro": _valor(fila, i["centro"]),
                    "fecha_inicio": _valor(fila, i["fini"]),
                    "fecha_fin": _valor(fila, i["ffin"]),
                    "duracion": _valor(fila, i["dur"]),
                    "turno": _valor(fila, i["turno"]),
                    "asignado_a": asignado,
                    "dni_asignado": _valor(fila, i["dni"]) or _valor(fila, i["dniaux"]),
                    "num_candidato": _valor(fila, i["numc"]),
                    "estado": estado,
                    "anotacion": _valor(fila, i["anot"]),
                    "hora": _valor(fila, i["hora"]),
                    "adjudicado": estado == ESTADO_ADJUDICADO,
                })
        finally:
            wb.close()
        return lista

    def _puesto_por_id(self, pid):
        pid = int(pid)
        for p in self.puestos():
            if p["id"] == pid:
                return p
        return None

    def estado(self):
        """Estado completo para la API (panel y pantalla)."""
        puestos = self.puestos()
        adjudicados = sum(1 for x in puestos if x["adjudicado"])

        # Enriquecer cada puesto asignado con el nombre de pila (sin apellidos) para la
        # pantalla de proyección. Se busca por nº de candidato; si no, primer token.
        por_numero = {str(c["numero"]): c["nombre_pila"] for c in self.candidatos()}
        for x in puestos:
            if x["asignado_a"]:
                x["nombre_pila"] = (por_numero.get(str(x["num_candidato"]))
                                    or x["asignado_a"].split()[0])
            else:
                x["nombre_pila"] = ""

        return {
            "puestos": puestos,
            "estados": self.config.get("estados", ESTADOS),
            "resumen": {
                "total": len(puestos),
                "adjudicados": adjudicados,
                "pendientes": len(puestos) - adjudicados,  # por ofertar
            },
            "actualizado": datetime.now().strftime("%H:%M:%S"),
            "ultimo_error": self.ultimo_error,
        }

    # ----------------------- Escritura (acciones) --------------------- #

    def asignar(self, puesto_id, numero_candidato, estado=None):
        """Asigna un candidato a un puesto y lo refleja en LOS DOS archivos."""
        cand = self.candidato_por_numero(numero_candidato)
        if cand is None:
            self.ultimo_error = f"No existe el candidato nº {numero_candidato}."
            return False
        estado = estado or ESTADO_ADJUDICADO
        with self.lock:
            puesto = self._puesto_por_id(puesto_id)
            if puesto is None:
                self.ultimo_error = "No se encontró el puesto indicado."
                return False
            anterior = puesto.get("num_candidato")

            # 1) Escribir en el archivo de PUESTOS.
            if not self._escribir_puesto(puesto_id, cand, estado):
                return False

            # 2) Si el puesto estaba asignado a otra persona, liberar a esa persona.
            if anterior and str(anterior) != str(cand["numero"]):
                self._limpiar_candidato(anterior)

            # 3) Escribir en el archivo de CANDIDATOS (constancia de su contrato).
            aceptada = self._texto_contrato(puesto)
            datos = f"Puesto #{int(puesto_id) - 1} · {datetime.now():%d/%m %H:%M}"
            if not self._marcar_candidato(cand, estado, aceptada, datos):
                self.ultimo_error = (
                    "El puesto se guardó, pero NO se pudo actualizar el listado de "
                    "candidatos (¿está abierto en Excel?). Ciérralo y vuelve a asignar."
                )
                return False

            self.ultimo_error = None
            return True

    def cambiar_estado(self, puesto_id, estado):
        with self.lock:
            puesto = self._puesto_por_id(puesto_id)
            if puesto is None:
                self.ultimo_error = "No se encontró el puesto indicado."
                return False
            if not self._editar_puesto(puesto_id, lambda ws, cab: (
                    _escribe(ws, int(puesto_id), cab, self.config["puestos"].get("col_estado"), estado),
                    _escribe(ws, int(puesto_id), cab, COL_HORA, datetime.now().strftime("%H:%M:%S")))):
                return False
            # Reflejar el estado también en el candidato asignado.
            if puesto.get("num_candidato"):
                self._marcar_candidato_estado(puesto["num_candidato"], estado)
            self.ultimo_error = None
            return True

    def liberar(self, puesto_id):
        """Quita la asignación del puesto (vuelve a 'por ofertar') y limpia al candidato."""
        with self.lock:
            puesto = self._puesto_por_id(puesto_id)
            if puesto is None:
                self.ultimo_error = "No se encontró el puesto indicado."
                return False
            anterior = puesto.get("num_candidato")
            p = self.config["puestos"]

            def acc(ws, cab):
                f = int(puesto_id)
                _escribe(ws, f, cab, p.get("col_nombre_acepta"), None)
                _escribe(ws, f, cab, p.get("col_dni_acepta"), None)
                _escribe(ws, f, cab, p.get("col_estado"), ESTADO_INICIAL)
                _escribe(ws, f, cab, COL_NUM, None)
                _escribe(ws, f, cab, COL_DNI_AUX, None)
                _escribe(ws, f, cab, COL_HORA, None)

            if not self._editar_puesto(puesto_id, acc):
                return False
            if anterior:
                self._limpiar_candidato(anterior)
            self.ultimo_error = None
            return True

    # ----------------------- Escritores internos ---------------------- #

    @staticmethod
    def _texto_contrato(puesto):
        partes = [puesto.get("centro", "")]
        if puesto.get("duracion"):
            partes.append(f"{puesto['duracion']} d")
        if puesto.get("turno"):
            partes.append(puesto["turno"])
        txt = " · ".join(p for p in partes if p)
        if puesto.get("fecha_inicio"):
            txt += f" ({puesto['fecha_inicio']}–{puesto.get('fecha_fin', '')})"
        return txt

    def _escribir_puesto(self, puesto_id, cand, estado):
        p = self.config["puestos"]
        tiene_col_dni = bool(_idx(self._cab_puestos(), p.get("col_dni_acepta")))

        def acc(ws, cab):
            f = int(puesto_id)
            _escribe(ws, f, cab, p.get("col_nombre_acepta"), cand["nombre"])
            _escribe(ws, f, cab, p.get("col_estado"), estado)
            _escribe(ws, f, cab, COL_NUM, cand["numero"])
            _escribe(ws, f, cab, COL_HORA, datetime.now().strftime("%H:%M:%S"))
            if tiene_col_dni:
                _escribe(ws, f, cab, p.get("col_dni_acepta"), cand["dni"])
            else:
                _escribe(ws, f, cab, COL_DNI_AUX, cand["dni"])

        return self._editar_puesto(puesto_id, acc)

    def _editar_puesto(self, puesto_id, accion):
        p = self.config["puestos"]
        return self._editar(p["ruta"], p["hoja"], accion,
                            extra_cols=(COL_NUM, COL_DNI_AUX, COL_HORA))

    def _marcar_candidato(self, cand, estado, aceptada, datos):
        c = self.config["candidatos"]

        def acc(ws, cab):
            f = cand["fila"]
            _escribe(ws, f, cab, c.get("col_oferta_estado"), estado)
            _escribe(ws, f, cab, c.get("col_oferta_aceptada"), aceptada)
            _escribe(ws, f, cab, c.get("col_oferta_datos"), datos)

        return self._editar(c["ruta"], c["hoja"], acc)

    def _marcar_candidato_estado(self, numero, estado):
        cand = self.candidato_por_numero(numero)
        if cand is None:
            return True
        c = self.config["candidatos"]
        return self._editar(c["ruta"], c["hoja"], lambda ws, cab:
                            _escribe(ws, cand["fila"], cab, c.get("col_oferta_estado"), estado))

    def _limpiar_candidato(self, numero):
        cand = self.candidato_por_numero(numero)
        if cand is None:
            return True
        c = self.config["candidatos"]

        def acc(ws, cab):
            f = cand["fila"]
            _escribe(ws, f, cab, c.get("col_oferta_estado"), None)
            _escribe(ws, f, cab, c.get("col_oferta_aceptada"), None)
            _escribe(ws, f, cab, c.get("col_oferta_datos"), None)

        return self._editar(c["ruta"], c["hoja"], acc)

    def _cab_puestos(self):
        p = self.config["puestos"]
        wb = load_workbook(p["ruta"], read_only=True)
        try:
            return _cabeceras(wb[p["hoja"]])
        finally:
            wb.close()

    # ----------------------- Edición segura de un Excel --------------- #

    def _editar(self, ruta, hoja, accion, extra_cols=()):
        """Abre el Excel, aplica `accion(ws, cabeceras)` y guarda de forma segura.

        Conserva fórmulas y formato (no usa data_only). Crea las columnas auxiliares que
        falten. Hace copia de seguridad y guardado atómico. Si el archivo está abierto en
        otro programa, deja `ultimo_error` y devuelve False (sin perder datos).
        """
        self._backup(ruta)
        wb = load_workbook(ruta)
        try:
            ws = wb[hoja]
            cab = list(_cabeceras(ws))
            for col in extra_cols:
                if _idx(cab, col) is None:
                    ws.cell(row=1, column=len(cab) + 1).value = col
                    cab.append(col)
            accion(ws, cab)
            self._guardar_atomico(wb, ruta)
            return True
        except PermissionError:
            self.ultimo_error = (
                f"No se pudo guardar «{os.path.basename(ruta)}»: está abierto en Excel. "
                "Ciérralo y vuelve a intentarlo."
            )
            return False
        finally:
            wb.close()

    @staticmethod
    def _backup(ruta):
        carpeta = os.path.join(os.path.dirname(os.path.abspath(ruta)), "backups")
        os.makedirs(carpeta, exist_ok=True)
        sello = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        try:
            shutil.copy2(ruta, os.path.join(carpeta, f"{sello}_{os.path.basename(ruta)}"))
        except OSError:
            pass  # una copia fallida no debe impedir trabajar

    @staticmethod
    def _guardar_atomico(wb, ruta):
        carpeta = os.path.dirname(os.path.abspath(ruta))
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=carpeta)
        os.close(fd)
        wb.save(tmp)
        os.replace(tmp, ruta)


# --------------------------------------------------------------------- #
#  Configuración (persistente junto al .exe o al script)
# --------------------------------------------------------------------- #

def _dir_persistente():
    if getattr(sys, "frozen", False):  # ejecutable PyInstaller
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


CONFIG_PATH = os.path.join(_dir_persistente(), "config.json")


def cargar_config():
    if os.path.exists(CONFIG_PATH):
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    return None


def guardar_config(cfg):
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
